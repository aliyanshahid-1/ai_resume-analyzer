from io import BytesIO
from pathlib import Path

import docx
import pypdf

SUPPORTED_TYPES = ["pdf", "docx"]


def _read_pdf(file_bytes: bytes) -> str:
    reader = pypdf.PdfReader(BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _read_docx(file_bytes: bytes) -> str:
    document = docx.Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                paragraphs.append(" | ".join(cells))

    return "\n".join(paragraphs)


def extract_resume_text(uploaded_file) -> str:
    """Extract readable text from a Streamlit-uploaded PDF or DOCX file."""
    suffix = Path(uploaded_file.name).suffix.lower()
    file_bytes = uploaded_file.getvalue()

    if suffix == ".pdf":
        text = _read_pdf(file_bytes)
    elif suffix == ".docx":
        text = _read_docx(file_bytes)
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")

    # Normalize excessive whitespace while preserving useful line breaks.
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines)
